"""Create the revision report with native SVG drawings and high-resolution PNG fallbacks.

Run with CODEX_PRIMARY_RUNTIME_PYTHON. Pass the two rendered plan PNGs as arguments.
"""
from pathlib import Path
import sys
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsmap
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
DEST = ROOT / 'docs' / 'MediLink-Revisi-Denah-Lantai-1-dan-2.docx'
doc = Document()
for border in doc.styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin = sec.bottom_margin = Mm(18)
sec.left_margin = sec.right_margin = Mm(18)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    st = doc.styles[name]
    st.font.name = 'DejaVu Sans'
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_after = Pt(7)
doc.styles['Normal'].font.size = Pt(10)
doc.styles['Normal'].paragraph_format.line_spacing = 1.12
doc.styles['Title'].font.size = Pt(24)
doc.styles['Title'].paragraph_format.space_after = Pt(10)
doc.styles['Heading 1'].font.size = Pt(15)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(13)
doc.styles['Heading 2'].font.size = Pt(12)

def para(text, style=None):
    p = doc.add_paragraph(text, style)
    return p

def table(headers, records, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c, w in zip(t.columns, widths):
        c.width = Mm(w)
    for row_index, record in enumerate([headers] + records):
        row = t.rows[0] if row_index == 0 else t.add_row()
        if row_index == 0:
            repeat = OxmlElement('w:tblHeader')
            row._tr.get_or_add_trPr().append(repeat)
        else:
            keep = OxmlElement('w:cantSplit')
            row._tr.get_or_add_trPr().append(keep)
        for i, (c, value, w) in enumerate(zip(row.cells, record, widths)):
            c.width = Mm(w)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr = c._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), '243F53' if row_index == 0 else ('F2F5F7' if row_index % 2 == 0 else 'FFFFFF'))
            tcPr.append(shade)
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:' + edge)
                for key, val in [('val','single'),('sz','5'),('color','D9D9D9')]:
                    e.set(qn('w:' + key), val)
                borders.append(e)
            tcPr.append(borders)
            mar = OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e = OxmlElement('w:' + edge)
                e.set(qn('w:w'), '90')
                e.set(qn('w:type'), 'dxa')
                mar.append(e)
            tcPr.append(mar)
            c.text = value
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.08
                for r in p.runs:
                    r.font.size = Pt(9)
                    r.font.bold = row_index == 0
                    r.font.color.rgb = RGBColor(255,255,255) if row_index == 0 else RGBColor(0,0,0)
    para('').paragraph_format.space_after = Pt(0)
    return t

para('Revisi ruang dan denah jaringan rumah sakit', 'Title')
para('MediLink versi 4.5   |   20 Oktober 2026', 'Subtitle')
para('Revisi ini memperbarui fungsi ruang pada kedua lantai, isi ruang, bukaan pintu dan instalasi jaringan. Denah lantai 1 dan 2 pada halaman berikutnya mengikuti data yang dipakai model 3D, sehingga nama ruang dan posisi perangkat dapat diperiksa bersama.')
para('Perubahan gedung utama lantai 1', 'Heading 1')
table(['Ruang sebelumnya', 'Fungsi terbaru', 'Isi ruang dan jaringan'], [
    ['Poliklinik', 'Laboratorium', 'Meja laboratorium, mikroskop, rak tabung, sink, komputer dan AP.'],
    ['Laboratorium', 'Depo farmasi', 'Kabinet obat, troli, komputer dan AP. Terhubung ke gudang dari dalam.'],
    ['Farmasi', 'Gudang farmasi', 'Kabinet stok obat, troli, komputer dan AP.'],
    ['Istirahat staf', 'Poli umum', 'Bed pemeriksaan, meja komputer, AP dan perlengkapan pemeriksaan.'],
    ['Kalibrasi alat medis', 'Poli THT', 'Kursi pemeriksaan THT, konsol alat, sink, komputer dan AP.'],
    ['Gudang alat medis', 'Poli gigi', 'Kursi dental, unit instrumen, sink, komputer dan AP.'],
], [43,43,88])
para('Akses farmasi dan pengunjung', 'Heading 1')
para('Pintu geser internal menghubungkan depo dan gudang farmasi melalui dinding bersama. Kabinet dan meja komputer ditempatkan di luar area pendekatan pintu. Kedua ruang tetap memiliki akses dari koridor utara.')
para('Satu AP outdoor dipasang pada tiang di taman pengunjung. Uplink LAN dan daya mengikuti tiang lalu jalur bawah tanah menuju gedung. Tablet tamu di area luar memakai VLAN 30; kebijakan simulasi menolak aksesnya ke server internal.')

doc.add_page_break()
para('Revisi back office dan lantai 2', 'Title')
para('Susunan ruang servis', 'Heading 1')
table(['Ruang sebelumnya', 'Fungsi terbaru', 'Perangkat dan akses'], [
    ['Dapur lama sisi barat', 'IPSRS dan kalibrasi', 'Dibagi oleh dinding menjadi dua ruang dengan pintu petugas masing masing.'],
    ['Bagian barat dapur lama', 'IPSRS bengkel pemeliharaan', 'Bench alat dan kabinet. Tanpa komputer dan AP; penerimaan alat dari barat.'],
    ['Bagian timur dapur lama', 'Kalibrasi alat medis', 'Bench instrumen dan kabinet. Tanpa komputer; satu AP.'],
    ['Administrasi servis', 'Gudang alat medis', 'Rak alat, terminal stok, switch servis, terminasi fiber dan panel daya.'],
    ['Linen bersih lama', 'Binatu', 'Mesin cuci komersial, meja stainless dan wadah linen. Penerimaan kotor dari belakang.'],
    ['Binatu lama', 'Linen bersih', 'Rak tekstil, meja lipat dan troli tertutup. Pintu penyerahan di timur.'],
    ['IPSRS lama', 'Dapur baru', 'Tepat di samping binatu, dipisahkan dinding. Sink, meja stainless, range, hood, fridge dan troli.'],
], [43,43,88])
para('Pintu bahan makanan dan pintu linen kotor terpisah di sisi belakang. Keduanya memiliki apron sendiri. Jalur petugas mengikuti koridor selatan, penghubung tertutup, lalu koridor back office. Dapur dan binatu tidak memiliki komputer kantor atau AP di dalam ruang.')
para('Area operasi lantai 2', 'Heading 1')
para('OK 1, OK 2 dan OK 3 tidak memiliki komputer maupun AP. Hanya satu AP dipertahankan di OK besar. Tidak ada komputer kantor di seluruh ruang operasi; peralatan tindakan dan lampu operasi tetap tersedia. Backbone fiber L2 tetap menuju nurse station dan distribusi sayap barat.')
para('Hasil pemeriksaan', 'Heading 1')
para('111 tes otomatis lulus. Seluruh 62 akses pintu dan portal, 84 rute masuk ruang, 163 rute sisi peralatan, delapan jalur servis serta delapan lintasan kamera roomtour lolos pemeriksaan. Audit tidak menemukan furnitur saling bertabrakan, keluar ruang, atau menghalangi akses pintu.')
para('Denah bersifat ilustratif tanpa skala. Ukuran memakai unit scene. Jalur kabel, radio, daya dan peralatan menggambarkan simulasi edukasi, bukan gambar kerja instalasi. Rendering WebGL pada GPU dan pengoperasian layar sentuh fisik belum diuji langsung.')

nsmap['asvg'] = 'http://schemas.microsoft.com/office/drawing/2016/SVG/main'
def add_plan(floor, png, width):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Mm(420), Mm(297)
    sec.top_margin = sec.bottom_margin = Mm(12)
    sec.left_margin = sec.right_margin = Mm(12)
    p = para('')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1
    shape = p.add_run().add_picture(str(png), width=Mm(width))
    shape._inline.docPr.set('descr', f'Denah ruang dan jaringan lantai {floor} MediLink versi 4.5 dengan legenda LAN fiber daya dan Wi Fi')
    svg = ROOT / 'assets' / f'denah-lantai-{floor}.svg'
    part = Part(PackURI(f'/word/media/denah-lantai-{floor}.svg'), 'image/svg+xml', svg.read_bytes(), doc.part.package)
    rid = doc.part.relate_to(part, RT.IMAGE)
    blip = shape._inline.xpath('.//a:blip')[0]
    extlst = OxmlElement('a:extLst')
    ext = OxmlElement('a:ext')
    ext.set('uri', '{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
    svgblip = OxmlElement('asvg:svgBlip')
    svgblip.set(qn('r:embed'), rid)
    ext.append(svgblip)
    extlst.append(ext)
    blip.append(extlst)

add_plan(1, Path(sys.argv[1]), 349)
add_plan(2, Path(sys.argv[2]), 396)
doc.core_properties.title = 'Revisi ruang dan denah jaringan rumah sakit'
doc.core_properties.subject = 'MediLink 4.5 denah lantai 1 dan 2'
doc.core_properties.author = 'MediLink'
doc.core_properties.keywords = 'MediLink, denah, jaringan, lantai 1, lantai 2, SVG'
doc.save(DEST)
print(DEST)
